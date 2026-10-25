"""🗓️ Class routine from the official Google Sheet.

The department routine sheet (config.json → routine.sheet_url) is downloaded as Excel,
the grid sheets ("CSE", "ECSE" …) are parsed, and for every section the bot keeps
ONE pinned, always-latest routine message in that section's 🗓️┃routine channel.
When the routine changes, the old message is deleted and replaced (no change-log posts).

Check: panel button "🔄 Check routine", or `!routine_check`. Auto-check every N hours.
"""
from __future__ import annotations

import io
import logging
import re
from itertools import zip_longest

import aiohttp
import discord
from discord.ext import commands, tasks

from core.config import DATA_DIR
from core.storage import append_audit, load_json, save_json

log = logging.getLogger("nubtk.routine")
STATE_FILE = DATA_DIR / "routine" / "state.json"
EVENTS_FILE = DATA_DIR / "events.json"

DAYS = ["Saturday", "Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
DAY_ALIASES = {d[:3].lower(): d for d in DAYS}
TIME_RE = re.compile(r"(\d{1,2})[.:](\d{2})\s*([AP]M)", re.I)
DATE_RE = re.compile(r"\b(\d{1,2}-[A-Za-z]{3}-\d{4})\b")


# ====================================================================== parsing (pure functions)
def _norm_time(label: str) -> str:
    """'03:30 PM – 04:30 AM' → '03:30 PM – 04:30 PM' (fixes the common AM/PM typo)."""
    found = TIME_RE.findall(label or "")
    if len(found) < 2:
        return (label or "").strip()
    (h1, m1, p1), (h2, m2, p2) = found[0], found[1]
    p1, p2 = p1.upper(), p2.upper()
    if p1 == "PM" and p2 == "AM" and int(h2) != 12:
        p2 = "PM"
    return f"{int(h1):02d}:{m1} {p1} – {int(h2):02d}:{m2} {p2}"


def _time_key(label: str) -> int:
    m = TIME_RE.search(label or "")
    if not m:
        return 9999
    h, mi, p = int(m.group(1)), int(m.group(2)), m.group(3).upper()
    if p == "PM" and h != 12:
        h += 12
    if p == "AM" and h == 12:
        h = 0
    return h * 60 + mi


def _cell(row, i) -> str:
    if i < len(row) and row[i] is not None:
        return str(row[i]).strip()
    return ""


def parse_grid(rows: list, sheet_name: str = "") -> list:
    """Parse a routine grid (Day | Room | Sec C.C. T | Sec C.C. T …, one or more blocks/shifts).

    Returns a list of dicts: day, time, room, section, course, teacher, sheet.
    """
    entries = []
    time_at: dict = {}        # column → time label   (from a row that contains time ranges)
    sec_cols: list = []       # columns where a "Sec" triple starts
    day_cols: list = []       # columns of "Day" (one per block)
    current_day: dict = {}    # block day col → current day name
    for raw in rows:
        row = [_cell(raw, i) for i in range(len(raw))]
        low = [c.lower() for c in row]
        # header row: Day, Room, Sec, C.C., T ...
        if "sec" in low and "day" in low:
            sec_cols = [i for i, c in enumerate(low) if c == "sec"]
            day_cols = [i for i, c in enumerate(low) if c == "day"]
            continue
        # time row: contains 2+ time ranges
        if sum(1 for c in row if len(TIME_RE.findall(c)) >= 2) >= 2:
            time_at = {i: _norm_time(c) for i, c in enumerate(row) if len(TIME_RE.findall(c)) >= 2}
            continue
        if not sec_cols or not day_cols:
            continue
        found = {dc: DAY_ALIASES.get(_cell(row, dc).lower()[:3]) for dc in day_cols}
        found = {dc: d for dc, d in found.items() if d}
        if found:
            # a new day starts: blocks whose own Day cell is empty follow the same day
            fallback = next(iter(found.values()))
            for dc in day_cols:
                current_day[dc] = found.get(dc, fallback)
        for sc in sec_cols:
            sec_raw, cc_raw, t_raw = _cell(row, sc), _cell(row, sc + 1), _cell(row, sc + 2)
            if not sec_raw or not cc_raw:
                continue
            block = max((dc for dc in day_cols if dc <= sc), default=None)
            if block is None or block not in current_day:
                continue
            room = _cell(row, block + 1)
            secs = [x.strip() for x in sec_raw.split("/")]
            ccs = [x.strip() for x in cc_raw.split("/")]
            ts = [x.strip() for x in t_raw.split("/")]
            last = ("", "", "")
            for s_, c_, t_ in zip_longest(secs, ccs, ts):
                s_ = s_ or last[0]
                c_ = c_ or last[1]
                t_ = t_ or last[2]
                last = (s_, c_, t_)
                if not s_ or not c_:
                    continue
                entries.append({
                    "day": current_day[block],
                    "time": time_at.get(sc, ""),
                    "room": room,
                    "section": s_.upper(),
                    "course": re.sub(r"\s+", " ", c_.upper()),
                    "teacher": t_.upper(),
                    "sheet": sheet_name,
                })
    return entries


def _fmt_date(v) -> str:
    """Cell value (datetime or text) → '24 Oct 2026' ('' if it is not a date)."""
    from datetime import date, datetime
    if isinstance(v, (datetime, date)):
        return v.strftime("%d %b %Y")
    txt = str(v or "").strip()
    m = DATE_RE.search(txt)
    if m:
        try:
            return datetime.strptime(m.group(1), "%d-%b-%Y").strftime("%d %b %Y")
        except ValueError:
            return m.group(1)
    m = re.search(r"\b(\d{4})-(\d{2})-(\d{2})\b", txt)
    if m:
        return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3))).strftime("%d %b %Y")
    return ""


def find_effective_date(rows: list) -> str:
    """Date next to (or under) the first 'Effective From' label at the top of the sheet."""
    top = [list(r) for r in rows[:10]]
    for ri, raw in enumerate(top):
        for ci, c in enumerate(raw):
            if "effective" not in str(c or "").lower():
                continue
            cands = [raw[j] for j in range(ci + 1, min(ci + 5, len(raw)))]
            if ri + 1 < len(top) and ci < len(top[ri + 1]):
                cands.insert(0, top[ri + 1][ci])  # value written in the cell below the label
            for v in [c, *cands]:
                d = _fmt_date(v)
                if d:
                    return d
    for raw in top:  # fallback: any date in the header
        for c in raw:
            d = _fmt_date(c)
            if d:
                return d
    return ""


CODE_RE = re.compile(r"^[A-Z]{2,5} ?\d{4}$")


def parse_titles(rows: list) -> dict:
    """'CSE 1205' → 'Digital Logic Design' from the Course Offerings sheet."""
    out = {}
    for raw in rows:
        row = [_cell(raw, i) for i in range(len(raw))]
        for i, c in enumerate(row[:-1]):
            cu = re.sub(r"\s+", " ", c.upper())
            nxt = row[i + 1]
            if CODE_RE.match(cu) and nxt and not CODE_RE.match(nxt.upper()) and len(nxt) > 3:
                out.setdefault(cu if " " in cu else re.sub(r"([A-Z]+)(\d+)", r"\1 \2", cu), nxt)
    return out


def section_routines(entries: list, settings, teacher: str, scope: str) -> dict:
    """section → sorted unique list of entries for that section."""
    codes = {info["code"].upper(): key for key, info in settings.courses.items()}
    out = {sec: [] for sec in settings.sections}
    seen = set()
    for e in entries:
        sec = settings.normalize_section(e["section"])
        if not sec:
            continue
        if scope != "all":
            if teacher and e["teacher"] != teacher.upper():
                continue
            if e["course"] not in codes:
                continue
        key = (sec, e["day"], e["time"], e["course"], e["room"], e["teacher"])
        if key in seen:
            continue
        seen.add(key)
        out[sec].append({k: e.get(k, "") for k in ("day", "time", "room", "course", "teacher", "sheet")})
    for sec in out:
        out[sec].sort(key=lambda e: (DAYS.index(e["day"]), _time_key(e["time"]), e["course"]))
    return out


def _line(e) -> str:
    return f"{e['day']} · {e['time']} · {e['course']} · Room {e['room']} · {e['teacher']}"


# ====================================================================== cog
class Routine(commands.Cog):
    def __init__(self, bot):
        self.bot = bot
        self.s = bot.settings
        cfg = self.s.raw.get("routine", {}) or {}
        self.sheet_url = cfg.get("sheet_url", "")
        self.sheet_names = cfg.get("sheets", ["CSE", "ECSE"])
        self.titles_sheet = cfg.get("titles_sheet", "Course Offerings")
        self.teacher = cfg.get("teacher_initial", "") or (self.s.raw.get("instructor", {}) or {}).get("initial", "")

        self.state: dict = load_json(STATE_FILE, {})
        self.titles = (self.state.get("_meta") or {}).get("titles", {}) or {}
        hours = float(cfg.get("auto_check_hours", 6) or 0)
        self.auto = tasks.loop(hours=hours)(self._auto) if hours > 0 else None
        brief = str(cfg.get("daily_brief_time", "07:30") or "")
        self.brief = None
        if re.fullmatch(r"\d{1,2}:\d{2}", brief):
            from datetime import time as dtime
            h, m = (int(x) for x in brief.split(":"))
            self.brief = tasks.loop(time=dtime(hour=h, minute=m, tzinfo=self.s.tz))(self._daily_brief)
        from datetime import time as _t
        self.midnight = tasks.loop(time=_t(hour=0, minute=1, tzinfo=self.s.tz))(self._midnight)

    async def cog_load(self):
        if self.auto:
            self.auto.start()
        if self.brief:
            self.brief.start()
        self.midnight.start()

    async def cog_unload(self):
        self.midnight.cancel()
        if self.auto:
            self.auto.cancel()
        if self.brief:
            self.brief.cancel()

    async def _auto(self):
        await self.bot.wait_until_ready()
        try:
            repost = self._meta().get("view_v") != self.VIEW_VERSION
            report, changed = await self.check(force_repost=repost, post_summary=False)
            if repost:
                self._meta()["view_v"] = self.VIEW_VERSION
                save_json(STATE_FILE, self.state)
        except Exception as e:  # noqa: BLE001
            log.warning("Auto routine check failed: %s", e)
            return
        if changed:
            await self.bot.send_log("🗓️ Routine changed", report, discord.Color.orange())

    # ------------------------------------------------------------ download + parse
    def _sheet_id(self) -> str:
        m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", self.sheet_url or "")
        return m.group(1) if m else ""

    async def fetch_rows(self) -> tuple:
        sid = self._sheet_id()
        if not sid:
            raise commands.BadArgument("No routine sheet link in config.json → `routine.sheet_url`.")
        url = f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx"
        async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=90)) as sess:
            async with sess.get(url, allow_redirects=True) as r:
                data = await r.read()
                ctype = r.headers.get("Content-Type", "")
        if r.status != 200 or "html" in ctype.lower() or not data.startswith(b"PK"):
            raise commands.BadArgument(
                "Could not download the routine sheet. In Google Sheets click **Share → General access → "
                "Anyone with the link (Viewer)**, then try again.")
        def _parse():  # heavy on the Tinker Board → run outside the event loop so buttons stay responsive
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            out, titles = {}, None
            for name in self.sheet_names:
                if name in wb.sheetnames:
                    out[name] = [list(r) for r in wb[name].iter_rows(values_only=True)]
            if self.titles_sheet in wb.sheetnames:
                titles = parse_titles([list(r) for r in wb[self.titles_sheet].iter_rows(values_only=True)])
            return out, titles, list(wb.sheetnames)

        import asyncio
        sheets, titles, names = await asyncio.to_thread(_parse)
        if titles:
            self.titles = titles
            self.state.setdefault("_meta", {})["titles"] = titles
        if not sheets:
            raise commands.BadArgument(f"Sheets {self.sheet_names} not found. Available: {', '.join(names)}")
        return sheets

    async def build(self):
        """→ ({section: [all classes of the section]}, {sheet name: effective date})"""
        sheets = await self.fetch_rows()
        import asyncio

        def _grid():
            entries, effective = [], {}
            for name, rows in sheets.items():
                entries += parse_grid(rows, name)
                effective[name] = find_effective_date(rows)
            return section_routines(entries, self.s, "", "all"), effective
        return await asyncio.to_thread(_grid)

    def mine(self, items: list) -> list:
        codes = {info["code"].upper() for info in self.s.courses.values()}
        t = (self.teacher or "").upper()
        return [e for e in items if (not t or e["teacher"] == t) and e["course"] in codes]

    # ------------------------------------------------------------ messages
    DAY_ICON = {"Saturday": "🟦", "Sunday": "🟩", "Monday": "🟨", "Tuesday": "🟧",
                "Wednesday": "🟥", "Thursday": "🟪", "Friday": "⬜"}

    def _course_name(self, code: str, short: bool = False) -> str:
        info = next((i for i in self.s.courses.values() if i["code"].upper() == code), {})
        if short and info.get("short"):
            return info["short"]
        return self.titles.get(code) or info.get("title") or ""

    def routine_embed(self, sec, items, effective, kind="all", updated="") -> discord.Embed:
        """kind = 'all' (full section routine) or 'mine' (only my classes). Easy to read on a phone."""
        s = self.s
        eve = " · Evening" if s.is_evening(sec) else ""
        t = (self.teacher or "").upper()
        if kind == "all":
            title = f"🗓️ Section {sec}{eve} — Class Routine"
            color = discord.Color(s.section_color(sec))
            head = (f"📅 **Effective from:** {effective or '—'}\n"
                    f"🕒 **Last updated:** {updated or s.now().strftime('%d %b %Y, %I:%M %p')}\n"
                    f"⭐ = class with {self.teacher or 'me'}")
        else:
            title = f"👨‍🏫 Section {sec} — classes with {self.teacher or 'me'}"
            color = discord.Color.gold()
            head = "Classes are on campus unless an online class is announced."
        e = discord.Embed(title=title, color=color, description=head)
        by_day: dict = {}
        for it in items:
            by_day.setdefault(it["day"], []).append(it)
        for day in DAYS:
            if day not in by_day:
                continue
            blocks = []
            for it in sorted(by_day[day], key=lambda x: _time_key(x["time"])):
                name = self._course_name(it["course"], short=True)
                lab = " 🧪" if "LAB" in (it["room"] or "").upper() or "lab" in name.lower() else ""
                star = "⭐ " if kind == "all" and t and it["teacher"] == t else ""
                first = f"{star}**{it['time']}**"
                if len(name) > 44:
                    name = name[:43].rstrip(" :-") + "…"
                second = f"{it['course']}" + (f" · {name}" if name else "") + lab
                room = str(it["room"] or "—")
                room = f"Room {room}" if room[:1].isdigit() else room
                third = room + (f" · 👨‍🏫 {it['teacher']}" if kind == "all" else "")
                blocks.append(f"{first}\n> {second}\n> 🏫 {third}")
            e.add_field(name=f"{self.DAY_ICON.get(day, '📅')} {day.upper()}", value="\n".join(blocks)[:1024], inline=False)
        if not items:
            e.add_field(name="No classes", value="Nothing found in the current routine.", inline=False)
        if kind == "all":
            e.set_footer(text="Official department routine · this pinned message is always the latest")
        return e

    VIEW_VERSION = 5  # bump when the buttons under the pinned routine change → auto re-post

    def _link_view(self):
        """Under the pinned routine: students' 📅 buttons + 📄 sheet link (teacher updates from the admin panel)."""
        v = discord.ui.View(timeout=None)
        v.add_item(discord.ui.Button(label="My section today (classes & deadlines)", emoji="📅", style=discord.ButtonStyle.primary,
                                     custom_id="nubtk:rt:today", row=0))
        v.add_item(discord.ui.Button(label="Tomorrow", emoji="📆", style=discord.ButtonStyle.secondary,
                                     custom_id="nubtk:rt:tomorrow", row=0))
        if self.sheet_url:
            v.add_item(discord.ui.Button(label="Open full routine sheet", emoji="📄", style=discord.ButtonStyle.link,
                                         url=self.sheet_url, row=1))
        v.stop()  # buttons are handled by on_interaction / the panel (work after restarts)
        return v

    # ------------------------------------------------------------ deadlines & tests (assignment, CT, quiz …)
    EVENT_TYPES = {"assignment": ("Assignment", "📝"), "ct": ("CT / Class test", "✍️"), "quiz": ("Quiz", "❓"),
                   "lab": ("Lab report", "🧪"), "presentation": ("Presentation", "🎤"), "exam": ("Exam", "📚"),
                   "other": ("Other", "📌")}

    def _events(self) -> dict:
        if not hasattr(self, "_ev"):
            self._ev = load_json(EVENTS_FILE, {})
        return self._ev

    def save_events(self):
        save_json(EVENTS_FILE, self._events())

    def add_event(self, kind, title, sections, when, details, by) -> str:
        import secrets
        eid = secrets.token_hex(3)
        self._events()[eid] = {"kind": kind, "title": title, "sections": sections, "when": when.isoformat(timespec="minutes"),
                               "details": details, "by": by, "created": self.s.now_str()}
        self.save_events()
        return eid

    def events_between(self, sec, start, end) -> list:
        """Events for a section (or all sections if sec is None) with start <= when < end."""
        from datetime import datetime as _dt
        out = []
        for eid, ev in self._events().items():
            w = _dt.fromisoformat(ev["when"])
            if not (start <= w < end):
                continue
            if sec and sec not in ev["sections"] and "all" not in ev["sections"]:
                continue
            out.append((w, eid, ev))
        out.sort(key=lambda x: x[0])
        return out

    def event_line(self, w, ev, show_sections=False) -> str:
        name, icon = self.EVENT_TYPES.get(ev["kind"], ("Other", "📌"))
        secs = f" · {', '.join(ev['sections'])}" if show_sections else ""
        det = f"\n   _{ev['details'][:120]}_" if ev.get("details") else ""
        return f"{icon} **{name}: {ev['title']}** — {w:%a %d %b, %I:%M %p}{secs}{det}"

    def _day_bounds(self, offset=0):
        from datetime import timedelta
        start = (self.s.now() + timedelta(days=offset)).replace(hour=0, minute=0, second=0, microsecond=0)
        return start, start + timedelta(days=1)

    def add_event_fields(self, e: discord.Embed, sec, offset=0, show_sections=False):
        from datetime import timedelta
        start, end = self._day_bounds(offset)
        today = self.events_between(sec, start, end)
        soon = self.events_between(sec, end, end + timedelta(days=7))
        if today:
            e.add_field(name="📌 Deadlines & tests on this day", inline=False,
                        value="\n".join(self.event_line(w, ev, show_sections) for w, _, ev in today)[:1024])
        if soon:
            e.add_field(name="⏳ Coming up (next 7 days)", inline=False,
                        value="\n".join(self.event_line(w, ev, show_sections) for w, _, ev in soon)[:1024])

    def prune_events(self):
        """Forget events older than 2 days."""
        from datetime import datetime as _dt, timedelta
        cutoff = self.s.now() - timedelta(days=2)
        ev = self._events()
        old = [k for k, v in ev.items() if _dt.fromisoformat(v["when"]) < cutoff]
        for k in old:
            ev.pop(k, None)
        if old:
            self.save_events()

    async def announce_event(self, eid):
        ev = self._events()[eid]
        from datetime import datetime as _dt
        w = _dt.fromisoformat(ev["when"])
        name, icon = self.EVENT_TYPES.get(ev["kind"], ("Other", "📌"))
        secs = list(self.s.sections) if "all" in ev["sections"] else ev["sections"]
        posted = []
        for sec in secs:
            ch = self.bot.channel_for(f"sec:{sec}:announcement")
            role = self.bot.role_for(f"role:sec:{sec}")
            if not ch:
                continue
            emb = discord.Embed(title=f"{icon} {name}: {ev['title']}", color=discord.Color.orange(),
                                description=(f"🗓️ **{w:%A, %d %B %Y}**\n⏰ **{w:%I:%M %p}**"
                                             + (f"\n\n{ev['details']}" if ev.get("details") else "")))
            emb.set_footer(text=f"{self.s.teacher} · Section {sec} · also shown under 📅 in 🗓️┃routine")
            try:
                await ch.send(content=role.mention if role else None, embed=emb)
                posted.append(sec)
            except discord.HTTPException:
                pass
        return posted

    # ------------------------------------------------------------ students: "my section's classes today"
    def _date_label(self, offset: int = 0) -> tuple:
        from datetime import timedelta
        d = self.s.now() + timedelta(days=offset)
        return d.strftime("%A"), d.strftime("%A, %d %B %Y")

    def section_day_embed(self, sec: str, offset: int = 0) -> discord.Embed:
        day, pretty = self._date_label(offset)
        items = [e for e in (self.state.get(sec) or {}).get("entries", []) if e["day"] == day]
        items.sort(key=lambda e: _time_key(e["time"]))
        t = (self.teacher or "").upper()
        word = "Today" if offset == 0 else "Tomorrow"
        e = discord.Embed(title=f"📅 Section {sec} — {word}",
                          description=f"🗓️ **{pretty}**\n", color=discord.Color(self.s.section_color(sec)))
        if not items:
            e.description += "\n🎉 No classes."
        else:
            e.description += "\n**🏫 Classes**"
            lines = []
            for it in items:
                title = self.titles.get(it["course"]) or next(
                    (i["title"] for i in self.s.courses.values() if i["code"].upper() == it["course"]), "")
                star = "⭐ " if it["teacher"] == t else ""
                lines.append(f"{star}🕘 **{it['time']}**\n   **{it['course']}** {title[:45]} · 🏫 Room **{it['room']}** · 👨‍🏫 {it['teacher']}")
            e.description += "\n" + "\n".join(lines)
        self.add_event_fields(e, sec, offset)
        e.set_footer(text="Only you can see this · from the official routine")
        return e

    def _section_of_channel(self, channel_id: int):
        for sec in self.s.sections:
            ch = self.bot.channel_for(f"sec:{sec}:routine")
            if ch and ch.id == channel_id:
                return sec
        return None

    def _section_of_member(self, member):
        rec = self.bot.store.get(member.id)
        return rec.get("section") if rec else None

    @commands.Cog.listener()
    async def on_interaction(self, interaction: discord.Interaction):
        if interaction.type != discord.InteractionType.component:
            return
        cid = (interaction.data or {}).get("custom_id", "")
        if not cid.startswith("nubtk:rt:"):
            return
        sec = self._section_of_channel(interaction.channel_id) or self._section_of_member(interaction.user)
        if not sec:
            await interaction.response.send_message("I couldn't find your section.", ephemeral=True)
            return
        offset = 1 if cid.endswith("tomorrow") else 0
        await interaction.response.send_message(embed=self.section_day_embed(sec, offset), ephemeral=True)

    async def _replace_pinned(self, sec, embeds) -> int | None:
        ch = self.bot.channel_for(f"sec:{sec}:routine")
        if ch is None:
            return None
        # clean channel: remove every old bot message (old routine, old change-logs, pin notices)
        try:
            async for m in ch.history(limit=200):
                if m.author.id == self.bot.user.id or m.type == discord.MessageType.pins_add:
                    try:
                        await m.delete()
                    except discord.HTTPException:
                        pass
        except discord.HTTPException:
            old = (self.state.get(sec) or {}).get("msg_id")
            if old:
                try:
                    await (await ch.fetch_message(int(old))).delete()
                except discord.HTTPException:
                    pass
        kw = {"embeds": embeds, "view": self._link_view()}
        msg = await ch.send(**kw)
        try:
            await msg.pin()
            async for m in ch.history(limit=5):
                if m.type == discord.MessageType.pins_add:
                    await m.delete()
        except discord.HTTPException:
            pass
        return msg.id

    # ------------------------------------------------------------ main check
    async def check(self, force_repost: bool = False, post_summary: bool = True):
        """Returns (report_text, any_changed)."""
        routines, eff_by_sheet = await self.build()
        lines, any_changed = [], False
        effs = {v for v in eff_by_sheet.values() if v}
        for sec, items in routines.items():
            sheets = [e.get("sheet") for e in items if e.get("sheet")]
            main_sheet = max(set(sheets), key=sheets.count) if sheets else ""
            effective = eff_by_sheet.get(main_sheet) or next(iter(effs), "")
            prev = self.state.get(sec) or {}
            old_set = {_line(e) for e in prev.get("entries", [])}
            new_set = {_line(e) for e in items}
            added, removed = sorted(new_set - old_set), sorted(old_set - new_set)
            first_time = "entries" not in prev
            changed = bool(added or removed) and not first_time
            eff_changed = prev.get("effective") != effective
            ch = self.bot.channel_for(f"sec:{sec}:routine")
            if ch is None:
                lines.append(f"• **{sec}** — ⚠️ no routine channel (run `!setup_server`)")
                continue
            updated = prev.get("updated") or ""
            if first_time or changed or eff_changed or not updated:
                updated = self.s.now().strftime("%d %b %Y, %I:%M %p")
            if first_time or changed or eff_changed or force_repost or not prev.get("msg_id"):
                mine = self.mine(items)
                embeds = [self.routine_embed(sec, items, effective, "all", updated)]
                if mine:
                    embeds.append(self.routine_embed(sec, mine, effective, "mine", updated))
                mid = await self._replace_pinned(sec, embeds)
                prev = {"entries": items, "effective": effective, "msg_id": mid, "updated": updated}
                self.state[sec] = prev
            if changed:
                any_changed = True
                append_audit(self.s, "routine_changed", section=sec, added=added, removed=removed)
                lines.append(f"• **{sec}** — 🔔 **updated** (+{len(added)} / −{len(removed)}) → {ch.mention}")
            elif first_time:
                lines.append(f"• **{sec}** — 📌 routine posted ({len(items)} classes, {len(self.mine(items))} yours) → {ch.mention}")
            else:
                lines.append(f"• **{sec}** — ✅ no change ({len(items)} classes, {len(self.mine(items))} yours)")
        save_json(STATE_FILE, self.state)
        report = f"🗓️ **Routine check** · effective from **{', '.join(sorted(effs)) or '—'}**\n" + "\n".join(lines)
        return report, any_changed

    # ------------------------------------------------------------ my classes (today / tomorrow / week)
    def my_classes(self, day: str) -> list:
        out = []
        for sec, st in self.state.items():
            if sec not in self.s.sections:
                continue
            for e in self.mine(st.get("entries", [])):
                if e["day"] == day:
                    out.append({**e, "section": sec})
        out.sort(key=lambda e: (_time_key(e["time"]), e["section"]))
        return out

    def day_embed(self, day: str, label: str) -> discord.Embed:
        items = self.my_classes(day)
        ours = {info["code"].upper(): info for info in self.s.courses.values()}
        from datetime import timedelta
        offset = next((i for i in range(7) if (self.s.now() + timedelta(days=i)).strftime("%A") == day), 0)
        pretty = (self.s.now() + timedelta(days=offset)).strftime("%A, %d %B %Y")
        e = discord.Embed(title=f"📅 {label} — classes, deadlines & tests", description=f"🗓️ **{pretty}**\n",
                          color=discord.Color.gold(), timestamp=self.s.now())
        if not items:
            e.description += "\n🎉 No classes."
            self.add_event_fields(e, None, offset, show_sections=True)
            return e
        lines = []
        for it in items:
            info = ours.get(it["course"], {})
            lab = " 🧪" if info.get("type") == "lab" else ""
            eve = " (Evening)" if self.s.is_evening(it["section"]) else ""
            lines.append(f"🕘 **{it['time']}**\n"
                         f"   **{it['course']}** {info.get('short', '')}{lab} · Section **{it['section']}**{eve} · 🏫 Room **{it['room']}**")
        e.description = (e.description + "\n" + "\n".join(lines))[:4000]
        self.add_event_fields(e, None, offset, show_sections=True)
        e.set_footer(text=f"{len(items)} class(es) · from the official routine")
        return e

    def week_embed(self) -> discord.Embed:
        e = discord.Embed(title="📆 My weekly routine", color=discord.Color.gold(), timestamp=self.s.now())
        ours = {info["code"].upper(): info for info in self.s.courses.values()}
        total = 0
        for day in DAYS:
            items = self.my_classes(day)
            if not items:
                continue
            total += len(items)
            e.add_field(name=f"📅 {day}", inline=False, value="\n".join(
                f"`{it['time']}` **{it['course']}** {ours.get(it['course'], {}).get('short', '')} · "
                f"{it['section']} · 🏫 {it['room']}" for it in items)[:1024])
        e.description = f"{total} classes per week" if total else "No classes found."
        return e

    async def _ensure_state(self):
        if not self.state:
            await self.check()

    def _today(self, offset: int = 0) -> str:
        from datetime import timedelta
        d = self.s.now() + timedelta(days=offset)
        return d.strftime("%A")

    def _meta(self) -> dict:
        return self.state.setdefault("_meta", {})

    async def _delete_old_brief(self):
        mid = self._meta().pop("brief_msg", None)
        ch = self.bot.channel_for("admin_channel")
        if mid and ch:
            try:
                await (await ch.fetch_message(int(mid))).delete()
            except discord.HTTPException:
                pass
        save_json(STATE_FILE, self.state)

    async def _midnight(self):
        await self.bot.wait_until_ready()
        await self._delete_old_brief()

    async def _daily_brief(self):
        await self.bot.wait_until_ready()
        await self._delete_old_brief()  # yesterday's message disappears
        ch = self.bot.channel_for("admin_channel")
        if not ch or not self.state:
            return
        day = self._today()
        self.prune_events()
        start, end = self._day_bounds(0)
        if not self.my_classes(day) and not self.events_between(None, start, end):
            return  # nothing today → no message
        try:
            msg = await ch.send(content="☀️ Good morning! Your classes today (this message is removed tomorrow):",
                                embed=self.day_embed(day, "Today"))
            self._meta()["brief_msg"] = msg.id
            save_json(STATE_FILE, self.state)
        except discord.HTTPException:
            pass

    @commands.command(name="deadlines")
    async def deadlines_cmd(self, ctx):
        from datetime import timedelta
        now = self.s.now()
        items = self.events_between(None, now - timedelta(hours=1), now + timedelta(days=60))
        if not items:
            await ctx.send("No upcoming deadlines or tests. Add one with 📌 Add deadline / test in the panel.")
            return
        await ctx.send(embed=discord.Embed(title="📌 Upcoming deadlines & tests", color=discord.Color.orange(),
                                           description="\n".join(f"`{eid}` " + self.event_line(w, ev, True)
                                                                  for w, eid, ev in items)[:4000]))

    @commands.command(name="today")
    async def today_cmd(self, ctx):
        await self._ensure_state()
        await ctx.send(embed=self.day_embed(self._today(), "Today"))

    @commands.command(name="tomorrow")
    async def tomorrow_cmd(self, ctx):
        await self._ensure_state()
        await ctx.send(embed=self.day_embed(self._today(1), "Tomorrow"))

    @commands.command(name="week")
    async def week_cmd(self, ctx):
        await self._ensure_state()
        await ctx.send(embed=self.week_embed())

    # ------------------------------------------------------------ commands
    @commands.command(name="routine_check")
    async def routine_check(self, ctx):
        msg = await ctx.send("⏳ Downloading the routine sheet…")
        report, _ = await self.check()
        await msg.edit(content=report[:1990])

    @commands.command(name="routine_repost")
    async def routine_repost(self, ctx):
        msg = await ctx.send("⏳ Re-posting routines…")
        report, _ = await self.check(force_repost=True)
        await msg.edit(content=report[:1990])

    @commands.command(name="routine")
    async def routine_show(self, ctx, section: str = "", which: str = "both"):
        """!routine · !routine 3D · !routine 3D all · !routine 3D mine"""
        sec = self.s.normalize_section(section) if section else None
        if section and not sec:
            if section.lower() in ("all", "mine", "both"):
                which, sec = section.lower(), None
            else:
                raise commands.BadArgument(f"Unknown section. Valid: {', '.join(self.s.sections)}")
        secs = [sec] if sec else list(self.s.sections)
        embeds = []
        for x in secs:
            st = self.state.get(x)
            if not st:
                continue
            items = st.get("entries", [])
            if which in ("both", "all"):
                embeds.append(self.routine_embed(x, items, st.get("effective", ""), "all", st.get("updated", "")))
            if which in ("both", "mine"):
                embeds.append(self.routine_embed(x, self.mine(items), st.get("effective", ""), "mine", st.get("updated", "")))
        if not embeds:
            raise commands.BadArgument("No routine saved yet — run `!routine_check` first.")
        for i in range(0, len(embeds), 4):
            await ctx.send(embeds=embeds[i:i + 4])


async def setup(bot):
    await bot.add_cog(Routine(bot))
