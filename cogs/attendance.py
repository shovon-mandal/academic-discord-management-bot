"""Voice-channel attendance with time tracking.

Start:  !att_start 2C [CSE1206] [minutes]   (or the ▶️ button in the admin panel)
        Every voice channel inside the 2C category counts.
Stop:   !att_stop 2C                         (or ⏹️ button, or automatic after <minutes>)
        → Excel file: Present/Absent, first join, last leave, minutes in class, % of class time
        → CSV copy (used by the semester report)
Report: !att_report 2C [course]              → Excel with one column per class + total %
"""
from __future__ import annotations

import csv
import logging
from collections import OrderedDict
from datetime import datetime, timedelta

import discord
from discord.ext import commands, tasks

from core.storage import ATTENDANCE_DIR, append_audit, load_json, save_json, write_csv

log = logging.getLogger("nubtk.attendance")
CSV_HEADER = ["SL", "Student ID", "Name", "Status", "First Join", "Last Leave", "Minutes", "Percent of Class",
              "Joins", "Discord Username"]


def _fmt_dur(seconds: float) -> str:
    seconds = int(max(0, seconds))
    h, r = divmod(seconds, 3600)
    m, s = divmod(r, 60)
    return f"{h}h {m:02d}m" if h else f"{m}m {s:02d}s"


class Attendance(commands.Cog):
    def __init__(self, bot):
        self.bot = bot
        self.s = bot.settings
        self.path = ATTENDANCE_DIR / "active_sessions.json"
        self.sessions: dict = load_json(self.path, {})
        self.min_percent = float(self.s.raw.get("attendance_min_percent", 0) or 0)

    async def cog_load(self):
        self.auto_stop.start()

    async def cog_unload(self):
        self.auto_stop.cancel()

    def _save(self):
        save_json(self.path, self.sessions)

    def _category(self, sec):
        return self.bot.channel_for(f"sec:{sec}:category")

    def _need_section(self, raw):
        sec = self.s.normalize_section(raw)
        if not sec:
            raise commands.BadArgument(f"Unknown section `{raw}`. Valid: {', '.join(self.s.sections)}")
        return sec

    # ------------------------------------------------------------ time tracking
    def _join(self, sec, member, channel_name, when: datetime | None = None):
        sess = self.sessions.get(sec)
        if not sess or member.bot or self.bot.is_admin(member):
            return
        when = when or self.s.now()
        p = sess["present"].setdefault(str(member.id), {
            "first": when.isoformat(timespec="seconds"), "last_leave": None, "total": 0.0,
            "open_since": None, "joins": 0, "channel": channel_name,
            "display": member.display_name, "username": member.name,
        })
        if p.get("open_since"):
            return
        p["open_since"] = when.isoformat(timespec="seconds")
        p["joins"] = p.get("joins", 0) + 1
        p["channel"] = channel_name
        self._save()

    def _leave(self, sec, member_id, when: datetime | None = None):
        sess = self.sessions.get(sec)
        if not sess:
            return
        p = sess["present"].get(str(member_id))
        if not p or not p.get("open_since"):
            return
        when = when or self.s.now()
        p["total"] = p.get("total", 0.0) + (when - datetime.fromisoformat(p["open_since"])).total_seconds()
        p["open_since"] = None
        p["last_leave"] = when.isoformat(timespec="seconds")
        self._save()

    # ------------------------------------------------------------ public API (used by commands + panel)
    async def start_session(self, sec: str, course: str | None, minutes: int, by: str):
        if sec in self.sessions:
            raise commands.BadArgument(f"Attendance for {sec} is already running. Stop it first.")
        course = course or self.s.default_course(sec)
        if course not in self.s.section_courses(sec):
            raise commands.BadArgument(f"{course} is not a course of {sec}.")
        now = self.s.now()
        self.sessions[sec] = {
            "section": sec, "course": course, "started_at": now.isoformat(timespec="seconds"),
            "ends_at": (now + timedelta(minutes=minutes)).isoformat(timespec="seconds") if minutes else None,
            "started_by": by, "present": {},
        }
        self._save()
        cat = self._category(sec)
        if cat:
            for vc in cat.voice_channels:
                for m in vc.members:
                    self._join(sec, m, vc.name, now)
        info = self.s.courses[course]
        ann = self.bot.channel_for(f"sec:{sec}:announcement")
        role = self.bot.role_for(f"role:sec:{sec}")
        if ann:
            try:
                await ann.send(f"{role.mention if role else ''} 🟢 **Attendance is open** for {info['code']} "
                               f"{info['short']} — join **🔊 voice channel** of this section now."
                               + (f" (closes in {minutes} min)" if minutes else ""))
            except discord.HTTPException:
                pass
        return (f"🟢 Attendance started: **{sec} · {info['code']} {info['short']}**"
                + (f" · auto-stop in **{minutes} min**" if minutes else " · no auto-stop")
                + f"\nAlready in voice: {len(self.sessions[sec]['present'])}")

    async def stop_session(self, sec: str):
        """Returns (summary_text, [discord.File...])."""
        sess = self.sessions.get(sec)
        if not sess:
            raise commands.BadArgument(f"No running attendance for {sec}.")
        end = self.s.now()
        for did in list(sess["present"]):
            self._leave(sec, did, end)
        self.sessions.pop(sec, None)
        self._save()
        summary, paths = self._finalize(sess, end)
        ann = self.bot.channel_for(f"sec:{sec}:announcement")
        if ann:
            try:
                await ann.send("🔴 Attendance is closed. Thank you!")
            except discord.HTTPException:
                pass
        return summary, [discord.File(p) for p in paths]

    # ------------------------------------------------------------ commands
    @commands.command(name="att_start")
    async def att_start(self, ctx, section: str, *args):
        sec = self._need_section(section)
        course, minutes = None, 0
        for a in args:
            if a.isdigit() and len(a) <= 3:
                minutes = int(a)
                continue
            c = self.s.normalize_course(a)
            if not c or c not in self.s.section_courses(sec):
                raise commands.BadArgument(f"`{a}` is not a course of {sec}: {', '.join(self.s.section_courses(sec))}")
            course = c
        await ctx.send(await self.start_session(sec, course, minutes, ctx.author.name))

    @commands.command(name="att_stop")
    async def att_stop(self, ctx, section: str):
        summary, files = await self.stop_session(self._need_section(section))
        await ctx.send(summary, files=files)

    @commands.command(name="att_status")
    async def att_status(self, ctx):
        await ctx.send(self.status_text())

    def status_text(self) -> str:
        if not self.sessions:
            return "No attendance session is running."
        now = self.s.now()
        lines = []
        for sec, x in self.sessions.items():
            live = sum(1 for p in x["present"].values() if p.get("open_since"))
            elapsed = (now - datetime.fromisoformat(x["started_at"])).total_seconds()
            lines.append(f"• **{sec}** {self.s.courses[x['course']]['code']} — running {_fmt_dur(elapsed)} · "
                         f"joined {len(x['present'])} · in voice now {live}"
                         + (f" · ends {x['ends_at'][11:16]}" if x.get("ends_at") else ""))
        return "\n".join(lines)

    @commands.command(name="att_report")
    async def att_report(self, ctx, section: str, course: str = None):
        text, files = self.build_report(self._need_section(section), course)
        await ctx.send(text, files=files)

    # ------------------------------------------------------------ finalize one session
    def _finalize(self, sess, end: datetime):
        sec, course = sess["section"], sess["course"]
        info = self.s.courses[course]
        started = datetime.fromisoformat(sess["started_at"])
        class_secs = max(1.0, (end - started).total_seconds())
        present = sess["present"]
        roster = self.bot.store.section_students(sec)
        roster_ids = {r["discord_id"] for r in roster}

        def row_for(p):
            pct = 100 * p.get("total", 0) / class_secs
            first = datetime.fromisoformat(p["first"]).strftime("%H:%M:%S")
            last = datetime.fromisoformat(p["last_leave"]).strftime("%H:%M:%S") if p.get("last_leave") else ""
            return first, last, round(p.get("total", 0) / 60, 1), round(min(pct, 100), 1), p.get("joins", 1)

        rows, n_present = [], 0
        for i, r in enumerate(roster, 1):
            p = present.get(r["discord_id"])
            if p:
                first, last, mins, pct, joins = row_for(p)
                status = "P" if pct >= self.min_percent else "Late/Short"
                n_present += status == "P"
                rows.append([i, r["student_id"], r["name"], status, first, last, mins, pct, joins, r.get("discord_name", "")])
            else:
                rows.append([i, r["student_id"], r["name"], "A", "", "", 0, 0, 0, r.get("discord_name", "")])
        for did, p in present.items():
            if did in roster_ids:
                continue
            rec = self.bot.store.get(did) or {}
            first, last, mins, pct, joins = row_for(p)
            rows.append([len(rows) + 1, rec.get("student_id", ""), rec.get("name", p["display"]),
                         "P* (other section)" if rec else "P* (unverified)", first, last, mins, pct, joins, p["username"]])

        folder = ATTENDANCE_DIR / f"{course}_{sec}"
        stem = f"{started:%Y-%m-%d_%H%M}"
        csv_path = folder / f"{stem}.csv"
        write_csv(csv_path, CSV_HEADER, rows)
        paths = [csv_path]
        meta = {
            "Course": f"{info['code']} — {info['title']}", "Section": sec + (" (Evening)" if self.s.is_evening(sec) else ""),
            "Date": f"{started:%d %b %Y (%A)}", "Started": f"{started:%H:%M:%S}", "Stopped": f"{end:%H:%M:%S}",
            "Class duration": _fmt_dur(class_secs), "Taken by": sess.get("started_by", ""),
            "Present": f"{n_present} / {len(roster)}",
        }
        try:
            xlsx = self._session_xlsx(folder / f"{stem}.xlsx", meta, rows)
            paths.insert(0, xlsx)
        except ImportError:
            log.warning("openpyxl missing — only CSV created")
        append_audit(self.s, "attendance", section=sec, course=course, present=n_present, total=len(roster),
                     file=csv_path.name, duration=int(class_secs))
        summary = (f"🗓️ **Attendance · {sec} · {info['code']} {info['short']}** · {started:%d %b %Y}\n"
                   f"⏱️ {started:%H:%M} → {end:%H:%M} (**{_fmt_dur(class_secs)}**)\n"
                   f"✅ Present **{n_present}/{len(roster)}** · ❌ Absent {sum(1 for r in rows if r[3] == 'A')}"
                   + (f" · Others {len(rows) - len(roster)}" if len(rows) > len(roster) else "")
                   + "\n📗 Open the Excel file to see each student's time in class.")
        return summary, paths

    def _session_xlsx(self, path, meta: dict, rows: list):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"
        ws.append(["Attendance Sheet"])
        ws["A1"].font = Font(bold=True, size=14)
        for k, v in meta.items():
            ws.append([k, v])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.append([])
        head_row = ws.max_row + 1
        ws.append(CSV_HEADER)
        for c in ws[head_row]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        green, red, yellow = (PatternFill("solid", fgColor=x) for x in ("C6EFCE", "FFC7CE", "FFEB9C"))
        for r in rows:
            ws.append(r)
            st = ws.cell(row=ws.max_row, column=4)
            st.fill = green if r[3] == "P" else red if r[3] == "A" else yellow
            st.alignment = Alignment(horizontal="center")
            ws.cell(row=ws.max_row, column=2).number_format = "@"
        for idx, w in enumerate([5, 14, 28, 16, 11, 11, 9, 10, 7, 22], 1):
            ws.column_dimensions[get_column_letter(idx)].width = w
        ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return path

    # ------------------------------------------------------------ semester report
    def build_report(self, sec: str, course: str | None):
        c = self.s.normalize_course(course) if course else self.s.default_course(sec)
        if not c or c not in self.s.section_courses(sec):
            raise commands.BadArgument(f"Course must be one of {', '.join(self.s.section_courses(sec))}")
        folder = ATTENDANCE_DIR / f"{c}_{sec}"
        files = sorted(folder.glob("*.csv")) if folder.exists() else []
        if not files:
            raise commands.BadArgument(f"No attendance taken yet for {c} {sec}.")
        students: dict = OrderedDict()
        for r in self.bot.store.section_students(sec):
            students[r["student_id"]] = {"name": r["name"], "cells": {}}
        dates = []
        for f in files:
            try:
                label = datetime.strptime(f.stem, "%Y-%m-%d_%H%M").strftime("%d %b %H:%M")
            except ValueError:
                label = f.stem
            dates.append(label)
            with open(f, encoding="utf-8-sig", newline="") as fh:
                for row in csv.DictReader(fh):
                    sid = row.get("Student ID", "")
                    if not sid:
                        continue
                    entry = students.setdefault(sid, {"name": row.get("Name", ""), "cells": {}})
                    status = row.get("Status", "")
                    mins = row.get("Minutes", "")
                    entry["cells"][label] = (status, mins)
        total = len(dates)
        report_rows = []
        for i, (sid, v) in enumerate(sorted(students.items()), 1):
            present = sum(1 for d in dates if v["cells"].get(d, ("",))[0].startswith("P"))
            pct = round(100 * present / total) if total else 0
            cells = []
            for d in dates:
                st, mins = v["cells"].get(d, ("A", ""))
                cells.append(f"{st[:1]} ({mins}m)" if st.startswith("P") and mins else (st[:1] or "A"))
            report_rows.append([i, sid, v["name"], *cells, present, total, pct])
        header = ["SL", "Student ID", "Name", *dates, "Present", "Total", "%"]
        out = []
        write_csv(ATTENDANCE_DIR / f"{c}_{sec}_report.csv", header, report_rows)
        try:
            out.append(discord.File(self._report_xlsx(ATTENDANCE_DIR / f"{c}_{sec}_report.xlsx", c, sec, header, report_rows)))
        except ImportError:
            out.append(discord.File(ATTENDANCE_DIR / f"{c}_{sec}_report.csv"))
        low = [r for r in report_rows if r[-1] < 60]
        text = (f"📈 **{self.s.courses[c]['code']} · {sec}** — {total} classes · {len(report_rows)} students · "
                f"below 60%: **{len(low)}**")
        return text, out

    def _report_xlsx(self, path, c, sec, header, rows):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = f"{c}_{sec}"[:31]
        ws.append([f"Attendance report · {self.s.courses[c]['code']} {self.s.courses[c]['title']} · Section {sec} · "
                   f"{self.s.semester}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
        ws.append(header)
        for cell in ws[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        red = PatternFill("solid", fgColor="FFC7CE")
        for r in rows:
            ws.append(r)
            ws.cell(row=ws.max_row, column=2).number_format = "@"
            if r[-1] < 60:
                ws.cell(row=ws.max_row, column=len(header)).fill = red
        for idx in range(1, len(header) + 1):
            ws.column_dimensions[get_column_letter(idx)].width = 5 if idx == 1 else 14 if idx == 2 else 26 if idx == 3 else 12
        ws.freeze_panes = "D4"
        wb.save(path)
        return path

    # ------------------------------------------------------------ auto stop + voice events
    @tasks.loop(seconds=30)
    async def auto_stop(self):
        now = self.s.now()
        for sec, sess in list(self.sessions.items()):
            if sess.get("ends_at") and now >= datetime.fromisoformat(sess["ends_at"]):
                try:
                    summary, files = await self.stop_session(sec)
                except Exception:
                    log.exception("auto-stop failed")
                    continue
                ch = self.bot.channel_for("admin_channel")
                if ch:
                    try:
                        await ch.send("⏱️ Auto-stopped.\n" + summary, files=files)
                    except discord.HTTPException:
                        pass

    @auto_stop.before_loop
    async def _before(self):
        await self.bot.wait_until_ready()

    @commands.Cog.listener()
    async def on_voice_state_update(self, member, before, after):
        if member.bot or not self.sessions:
            return
        b_id = before.channel.id if before.channel else None
        a_id = after.channel.id if after.channel else None
        if b_id == a_id:
            return
        for sec in list(self.sessions):
            cat = self._category(sec)
            if not cat:
                continue
            was_in = before.channel is not None and before.channel.category_id == cat.id
            now_in = after.channel is not None and after.channel.category_id == cat.id
            if was_in and not now_in:
                self._leave(sec, member.id)
            elif now_in and not was_in:
                self._join(sec, member, after.channel.name)


async def setup(bot):
    await bot.add_cog(Attendance(bot))
