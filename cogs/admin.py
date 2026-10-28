"""Admin commands (work only in the admin / bot-logs channel, only for admins)."""
from __future__ import annotations

import asyncio
import io
import logging
import re
from datetime import timedelta

import discord
from discord.ext import commands

from core.config import DATA_DIR
from core.parsing import clean_name
from core.storage import (STUDENTS_DIR, append_audit, load_roster, parse_roster_text, save_roster)
from core.system import zip_files

log = logging.getLogger("nubtk.admin")
GREEN, RED, GOLD, BLUE = discord.Color.green(), discord.Color.red(), discord.Color.gold(), discord.Color.blurple()

# (command, what it does, example)
HELP = [
    ("🏗️ Setup", [
        ("panel", "Re-post the button control panel", "!panel"),
        ("setup_server", "Create / sync all channels & roles (asks to confirm)", "!setup_server"),
        ("setup_welcome", "Post the welcome buttons + instructor cards AGAIN in every channel (only when you want; text/link changes update silently)", "!setup_welcome"),
        ("invite", "Invite link + join guide for students", "!invite"),
        ("reset_server", "⚠️ Delete EVERYTHING, build fresh (asks to confirm)", "!reset_server"),
    ]),
    ("🎓 Students", [
        ("stats", "Count per section & course", "!stats"),
        ("list <sec>", "All students of a section", "!list 3D"),
        ("find <id/name>", "Search a student", "!find 11250121964  ·  !find Rahim"),
        ("verify @user <sec> <id> <name>", "Verify manually", "!verify @rahim 3D 11250121964 Rahim Uddin"),
        ("unverify @user", "Remove access + record", "!unverify @rahim"),
        ("move @user <sec>", "Change section", "!move @rahim 3F"),
        ("rename @user <name>", "Fix name / nickname", "!rename @rahim Md Rahim Uddin"),
        ("clean_welcome [all]", "Clean welcome + ID channels: verified/left students (auto) · `all` = also students still waiting", "!clean_welcome all"),
        ("add_section @user <sec…>", "Give access to an EXTRA section (student in 3D and 4B)", "!add_section @rahim 4B"),
        ("remove_section @user <sec…>", "Remove an extra section", "!remove_section @rahim 4B"),
        ("pending", "Requests waiting for approval", "!pending"),
        ("notify_unverified", "Remind unverified members", "!notify_unverified"),
        ("kick_unverified <days> [confirm]", "List / kick never-verified", "!kick_unverified 7  →  !kick_unverified 7 confirm"),
        ("roster_upload <sec>", "Upload official list (attach CSV)", "!roster_upload 3D  + attach 3D.csv"),
        ("missing <sec>", "Roster students not joined yet", "!missing 3D"),
    ]),
    ("📢 Communication", [
        ("announce <target> <text>", "Post + ping (all / 3D / 3D:lab / CSE2109)",
         "!announce 3D Class test on Sunday at 10 AM"),
        ("announce_silent <target> <text>", "Same, no ping", "!announce_silent all Slides are uploaded"),
        ("dm <target> <text>", "DM every student", "!dm 4P Tomorrow's class is online"),
        ("marks_dm <title>", "Send each student their own marks (attach CSV/Excel)", "!marks_dm CT-1 Marks  + attach ct1.xlsx"),
        ("scheduled", "List / cancel scheduled announcements", "!scheduled"),
    ]),
    ("🗓️ Routine & reminders", [
        ("today", "My day: classes (section, time, room) + deadlines & tests", "!today"),
        ("deadlines", "Upcoming deadlines & tests (add them with 📌 in the panel)", "!deadlines"),
        ("tomorrow", "My classes tomorrow", "!tomorrow"),
        ("week", "My whole weekly routine", "!week"),
        ("routine_check", "Update routine: check the online sheet, update pinned routines, announce changes", "!routine_check"),
        ("routine [sec] [all|mine]", "Show saved routine: full section / only my classes", "!routine 3D  ·  !routine 3D mine"),
        ("routine_repost", "Re-post every section's pinned routine", "!routine_repost"),
        ("(right-click a message → Apps → ⏰ Set reminder)", "Re-announce a message now or later", "deadline / CT reminder"),
    ]),
    ("🗓️ Attendance", [
        ("att_start <sec> [course] [minutes]", "Start voice attendance", "!att_start 3D 15  ·  !att_start 3E CSE2110 30"),
        ("att_stop <sec>", "Stop + get CSV", "!att_stop 3D"),
        ("att_status", "Running sessions", "!att_status"),
        ("att_report <sec> [course]", "Semester % per student", "!att_report 3D CSE2109"),
    ]),
    ("💾 Data & System", [
        ("excel [target]", "Student list as Excel", "!excel  ·  !excel 3D  ·  !excel CSE2217"),
        ("export [target]", "Student CSV/JSON files", "!export CSE2109 3D"),
        ("backup", "Zip of all data", "!backup"),
        ("status", "Uptime, CPU temp, RAM", "!status"),
        ("logs [lines]", "Last log lines", "!logs 50"),
        ("update", "Pull from GitHub + restart", "!update"),
        ("restart", "Restart the bot", "!restart"),
    ]),
]


class Admin(commands.Cog):
    def __init__(self, bot):
        self.bot = bot
        self.s = bot.settings

    # ------------------------------------------------------------ helpers
    @property
    def ver(self):
        return self.bot.get_cog("Verification")

    def help_embeds(self) -> list:
        """One embed per group; every command shows an example."""
        p = self.s.prefix
        out = []
        colors = [BLUE, GREEN, GOLD, discord.Color.teal(), discord.Color.dark_grey()]
        for i, (title, cmds) in enumerate(HELP):
            body = "\n".join((f"**{c}** — {d}" if c.startswith("(") else f"**`{p}{c}`** — {d}\n↳ e.g. `{ex}`")
                              for c, d, ex in cmds)
            e = discord.Embed(title=title, description=body[:4000], color=colors[i % len(colors)])
            out.append(e)
        out[0].set_author(name="🛠️ Admin commands — only you can see this channel")
        return out

    def help_embed(self) -> discord.Embed:
        return self.help_embeds()[0]

    async def send_long(self, ctx, text: str, filename: str | None = None):
        if len(text) > 6000 and filename:
            await ctx.send(file=discord.File(io.BytesIO(text.encode("utf-8")), filename=filename))
            return
        chunk = ""
        for line in text.splitlines():
            if len(chunk) + len(line) + 1 > 1900:
                await ctx.send(f"```\n{chunk}```")
                chunk = ""
            chunk += line + "\n"
        if chunk.strip():
            await ctx.send(f"```\n{chunk}```")

    def need_section(self, raw) -> str:
        sec = self.s.normalize_section(raw)
        if not sec:
            raise commands.BadArgument(f"Unknown section `{raw}`. Valid: {', '.join(self.s.sections)}")
        return sec

    def resolve_target(self, raw: str):
        t = raw.strip()
        if t.lower() in ("all", "everyone", "sobai"):
            return "all", None
        lab = False
        m = re.match(r"(?i)^(.+?)[:\-_ ]?lab$", t)
        if m and self.s.normalize_section(m.group(1)):
            lab, t = True, m.group(1)
        sec = self.s.normalize_section(t)
        if sec:
            return ("section_lab" if lab else "section"), sec
        course = self.s.normalize_course(t)
        if course:
            return "course", course
        return None, None

    def target_info(self, kind, val):
        """-> (channel, roles, recipients, label)"""
        s, b = self.s, self.bot
        if kind == "all":
            return (b.channel_for("all:notice"), [b.role_for("role:verified")], b.store.all_active(), "All students")
        if kind in ("section", "section_lab"):
            ch = None
            if kind == "section_lab":
                ch = b.channel_for(f"sec:{val}:lab")
            ch = ch or b.channel_for(f"sec:{val}:announcement")
            return ch, [b.role_for(f"role:sec:{val}")], b.store.section_students(val), f"Section {val}"
        theory = s.theory_course_for(val)
        secs = s.course_sections(val)
        return (b.channel_for(f"group:{theory}:announcement") or b.channel_for(f"sec:{secs[0]}:announcement"),
                [b.role_for(f"role:sec:{x}") for x in secs], b.store.course_students(val),
                f"{s.courses[val]['code']} ({', '.join(secs)})")

    # ------------------------------------------------------------ invite
    @commands.command(name="invite")
    async def invite(self, ctx):
        wc = self.bot.channel_for("welcome_channel")
        if wc is None:
            raise commands.BadArgument("Run `!setup_server` first.")
        url = await self.bot.get_invite_url()
        if not url:
            raise commands.BadArgument("Could not create an invite link.")
        inv = type("Inv", (), {"url": url})
        ex = list(self.s.sections)[0]
        guide = (
            f"📢 {self.s.semester} · {self.s.teacher} — Join our class Discord server\n\n"
            f"1️⃣ Click this link: {inv.url}\n"
            "   (No Discord account? Create one first at discord.com or in the Discord app.)\n"
            "2️⃣ In the server, open the welcome channel and click your section button.\n"
            f"3️⃣ Write your Student ID ({self.s.id_length} digits) and full official name, then click Submit.\n"
            "✅ Done! Your class channels will open right away.\n\n"
            f"Button not working? In the id-name-section channel, type: {ex} 11250121964 Your Full Name\n\n"
            "Sections:\n" + "\n".join(f"• {x} — {self.s.short_for_section(x)}"
                                        + (" (Evening)" if self.s.is_evening(x) else "") for x in self.s.sections)
        )
        await ctx.send(f"🔗 Permanent invite: {inv.url}\n\n👇 Copy the message below and post it in your Facebook / WhatsApp class group:")
        await ctx.send(f"```\n{guide}\n```")

    # ------------------------------------------------------------ help / stats
    @commands.command(name="help", aliases=["commands"])
    async def help_cmd(self, ctx):
        await ctx.send(embeds=self.help_embeds())

    @commands.command(name="stats")
    async def stats(self, ctx):
        s, store = self.s, self.bot.store
        e = discord.Embed(title="📊 Students", color=GREEN, timestamp=s.now())
        total = 0
        for sec in s.sections:
            n = len(store.section_students(sec))
            total += n
            e.add_field(name=f"{sec}{' (Eve)' if s.is_evening(sec) else ''}",
                        value=f"**{n}** · {s.codes_label(s.section_courses(sec))}", inline=True)
        g = ctx.guild
        humans = [m for m in g.members if not m.bot]
        unverified = [m for m in humans if not self.ver.is_verified(m) and not self.bot.is_admin(m)]
        course_lines = []
        for c in s.courses:
            if s.course_sections(c):
                course_lines.append(f"{s.courses[c]['code']} {s.courses[c]['short']}: **{len(store.course_students(c))}**")
        e.add_field(name="Courses", value="\n".join(course_lines) or "-", inline=False)
        e.description = (f"Verified: **{total}** · Unverified members: **{len(unverified)}** · "
                         f"Pending approval: **{len(self.bot.pending.items)}** · Server members: {len(humans)}")
        await ctx.send(embed=e)

    @commands.command(name="list")
    async def list_cmd(self, ctx, section: str):
        sec = self.need_section(section)
        rows = self.bot.store.section_students(sec)
        if not rows:
            await ctx.send(f"No verified students in {sec} yet.")
            return
        lines = [f"Section {sec} · {self.s.section_courses_text(sec)} · {len(rows)} students", ""]
        lines += [f"{i:>3}. {r['student_id']}  {r['name'][:28]:<28} @{r.get('discord_name', '')}"
                  + ("" if r.get("section") == sec else f"  (extra, main {r.get('section')})")
                  for i, r in enumerate(rows, 1)]
        await self.send_long(ctx, "\n".join(lines), filename=f"{sec}_list.txt")

    @commands.command(name="find")
    async def find(self, ctx, *, query: str):
        hits = self.bot.store.search(query)[:20]
        if not hits:
            await ctx.send("Nothing found.")
            return
        lines = [f"{r['section']:<3} {r['student_id']}  {r['name'][:26]:<26} @{r.get('discord_name', '')} "
                 f"[{r.get('status', 'active')}] <{r['discord_id']}>" for r in hits]
        await self.send_long(ctx, "\n".join(lines))

    # ------------------------------------------------------------ manual management
    @commands.command(name="verify")
    async def verify(self, ctx, member: discord.Member, section: str, student_id: str, *, name: str):
        sec = self.need_section(section)
        sid = re.sub(r"\D", "", student_id)
        if len(sid) != self.s.id_length:
            raise commands.BadArgument(f"Student ID must be {self.s.id_length} digits.")
        ok, text = await self.ver.apply_verification(ctx.guild, member, sec, sid, clean_name(name) or name.strip(),
                                                     by=f"admin:{ctx.author.name}")
        p = self.bot.pending.find_by_member(member.id)
        if ok and p:
            self.bot.pending.pop(p[0])
        await ctx.send(("✅ " if ok else "❌ ") + text)

    @commands.command(name="unverify")
    async def unverify(self, ctx, member: discord.Member):
        rec = await self.ver.remove_verification(ctx.guild, member, by=ctx.author.name)
        await ctx.send(f"✅ {member.mention} unverified." + (f" (was {rec['section']} · {rec['student_id']})" if rec else ""))

    @commands.command(name="move")
    async def move(self, ctx, member: discord.Member, section: str):
        sec = self.need_section(section)
        rec = self.bot.store.get(member.id)
        if not rec:
            raise commands.BadArgument("This member is not verified. Use `!verify` instead.")
        ok, text = await self.ver.apply_verification(ctx.guild, member, sec, rec["student_id"], rec["name"],
                                                     by=f"admin-move:{ctx.author.name}")
        await ctx.send(("✅ Moved. " if ok else "❌ ") + text)

    @commands.command(name="clean_welcome")
    async def clean_welcome(self, ctx, mode: str = ""):
        everything = mode.lower() in ("all", "everything")
        n = await self.ver.sweep_join_channels(everything=everything)
        who = "all students' messages" if everything else "verified / left students"
        await ctx.send(f"🧹 Cleaned #welcome and #id-name-section: **{n}** message(s) removed ({who}). "
                       "Pinned instructions were kept.")

    @commands.command(name="add_section")
    async def add_section(self, ctx, member: discord.Member, *sections: str):
        secs = [self.need_section(x) for x in sections]
        if not secs:
            raise commands.BadArgument("Write at least one section, e.g. `!add_section @rahim 4B`.")
        ok, text = await self.ver.add_extra_sections(ctx.guild, member, secs, by=f"admin:{ctx.author.name}")
        await ctx.send(("✅ " if ok else "❌ ") + text)

    @commands.command(name="remove_section")
    async def remove_section(self, ctx, member: discord.Member, *sections: str):
        secs = [self.need_section(x) for x in sections]
        ok, text = await self.ver.remove_extra_sections(ctx.guild, member, secs, by=f"admin:{ctx.author.name}")
        await ctx.send(("✅ " if ok else "❌ ") + text)

    @commands.command(name="rename")
    async def rename(self, ctx, member: discord.Member, *, name: str):
        rec = self.bot.store.get(member.id)
        if not rec:
            raise commands.BadArgument("This member is not verified.")
        ok, text = await self.ver.apply_verification(ctx.guild, member, rec["section"], rec["student_id"], clean_name(name) or name.strip(),
                                                     by=f"admin-rename:{ctx.author.name}")
        await ctx.send(("✅ " if ok else "❌ ") + text)

    @commands.command(name="pending")
    async def pending(self, ctx):
        items = self.bot.pending.items
        if not items:
            await ctx.send("✅ No pending requests.")
            return
        for pid, it in list(items.items())[:15]:
            e = discord.Embed(title=f"🟡 Pending · {it['section']} · {it['student_id']}", color=GOLD,
                              description=f"<@{it['discord_id']}> · {it['name']}\n" + "\n".join(f"• {r}" for r in it["reasons"]))
            e.set_footer(text=f"pending id {pid} · {it['created_at']}")
            await ctx.send(embed=e, view=self.ver.pending_view(pid))

    @commands.command(name="notify_unverified")
    async def notify_unverified(self, ctx):
        msg = await ctx.send("⏳ Checking ID channel first...")
        v, e = await self.ver.scan_id_channel(reply_errors=True)
        idc = self.bot.channel_for("id_channel")
        wc = self.bot.channel_for("welcome_channel")
        targets = [m for m in ctx.guild.members if not m.bot and not self.bot.is_admin(m)
                   and not self.ver.is_verified(m) and not self.bot.pending.find_by_member(m.id)]
        sent = failed = 0
        embed = discord.Embed(
            title="🔔 Verification reminder",
            description=(f"You are not verified on **{ctx.guild.name}** yet.\n"
                         f"Go to {wc.mention if wc else '#welcome'} → click your section → fill the form.\n"
                         f"Or type in {idc.mention if idc else '#id-name-section'}: "
                         f"`{list(self.s.sections)[0]} 11250121964 Your Full Name`"),
            color=GOLD)
        for m in targets:
            try:
                await m.send(embed=embed)
                sent += 1
            except discord.HTTPException:
                failed += 1
            await asyncio.sleep(0.7)
        if idc and targets:
            for i in range(0, len(targets), 40):
                chunk = " ".join(m.mention for m in targets[i:i + 40])
                try:
                    await idc.send(f"👋 {chunk}\nYou are not verified yet — click your section button in "
                                   f"{wc.mention if wc else '#welcome'} or type here: "
                                   f"`SECTION STUDENT_ID FULL_NAME`")
                except discord.HTTPException:
                    pass
        await msg.edit(content=(f"✅ Done.\n• Verified from existing messages: **{v}**\n"
                                f"• Messages with errors (replied): **{e}**\n• Unverified now: **{len(targets)}**\n"
                                f"• DM sent: **{sent}** · DM failed: **{failed}**"))

    @commands.command(name="kick_unverified")
    async def kick_unverified(self, ctx, days: int, confirm: str = ""):
        cutoff = discord.utils.utcnow() - timedelta(days=days)
        targets = [m for m in ctx.guild.members if not m.bot and not self.bot.is_admin(m)
                   and not self.ver.is_verified(m) and m.joined_at and m.joined_at < cutoff]
        if not targets:
            await ctx.send(f"No unverified member older than {days} days.")
            return
        if confirm.lower() != "confirm":
            names = "\n".join(f"{m.name} (joined {m.joined_at:%Y-%m-%d})" for m in targets[:60])
            await self.send_long(ctx, f"{len(targets)} member(s) would be kicked:\n{names}")
            await ctx.send(f"Run `{self.s.prefix}kick_unverified {days} confirm` to kick them.")
            return
        kicked = 0
        for m in targets:
            try:
                await m.kick(reason=f"Unverified for {days}+ days")
                kicked += 1
            except discord.HTTPException:
                pass
        append_audit(self.s, "kick_unverified", days=days, count=kicked, by=ctx.author.name)
        await ctx.send(f"✅ Kicked {kicked} member(s).")

    # ------------------------------------------------------------ roster
    @commands.command(name="roster_upload")
    async def roster_upload(self, ctx, section: str):
        sec = self.need_section(section)
        if not ctx.message.attachments:
            raise commands.BadArgument("Attach a .csv / .txt file (one student per line: ID, Name).")
        raw = await ctx.message.attachments[0].read()
        entries = parse_roster_text(raw.decode("utf-8-sig", errors="replace"), self.s.id_length)
        if not entries:
            raise commands.BadArgument("No valid student IDs found in the file.")
        path = save_roster(sec, entries)
        note = "" if self.s.roster_check else ("\nℹ️ `roster_check` is **false** in config.json — the roster is "
                                               "only used by `!missing`. Set it to true to block wrong IDs.")
        await ctx.send(f"✅ Roster for **{sec}** saved: **{len(entries)}** students → `{path.name}`{note}")

    @commands.command(name="missing")
    async def missing(self, ctx, section: str):
        sec = self.need_section(section)
        roster = load_roster(sec)
        if not roster:
            raise commands.BadArgument(f"No roster for {sec}. Use `!roster_upload {sec}` with a CSV.")
        verified_ids = {r["student_id"] for r in self.bot.store.section_students(sec)}
        missing = sorted(set(roster) - verified_ids)
        extra = sorted(verified_ids - set(roster))
        lines = [f"Section {sec}: roster {len(roster)} · verified {len(verified_ids)} · missing {len(missing)}", "",
                 "NOT JOINED / NOT VERIFIED:"] + [f"  {sid}  {roster[sid]}" for sid in missing]
        if extra:
            lines += ["", "VERIFIED BUT NOT IN ROSTER:"] + [f"  {sid}" for sid in extra]
        await self.send_long(ctx, "\n".join(lines), filename=f"{sec}_missing.txt")

    # ------------------------------------------------------------ communication
    async def _announce(self, ctx, target: str, text: str, ping: bool, title: str = "📢 Announcement"):
        kind, val = self.resolve_target(target)
        if not kind:
            raise commands.BadArgument("Target must be `all`, a section like `3D`, `3D:lab`, or a course like `CSE2109`.")
        ch, roles, _, label = self.target_info(kind, val)
        if ch is None:
            raise commands.BadArgument("Announcement channel not found. Run `!setup_server`.")
        if not text and not ctx.message.attachments:
            raise commands.BadArgument("Write the announcement text after the target.")
        files = [await a.to_file() for a in ctx.message.attachments[:10]]
        embed = discord.Embed(title=title or "📢 Announcement", description=text or None, color=GOLD,
                              timestamp=self.s.now())
        embed.set_footer(text=f"{self.s.teacher} · {label}")
        content = " ".join(r.mention for r in roles if r) if ping else None
        kwargs = {"embed": embed}
        if content:
            kwargs["content"] = content
        if files:
            kwargs["files"] = files
        msg = await ch.send(**kwargs)
        await ctx.send(f"✅ Posted in {ch.mention} → {msg.jump_url}")

    @commands.command(name="announce")
    async def announce(self, ctx, target: str, *, text: str = ""):
        await self._announce(ctx, target, text, ping=True)

    @commands.command(name="announce_silent")
    async def announce_silent(self, ctx, target: str, *, text: str = ""):
        await self._announce(ctx, target, text, ping=False)

    @commands.command(name="dm")
    async def dm(self, ctx, target: str, *, text: str):
        kind, val = self.resolve_target(target)
        if not kind:
            raise commands.BadArgument("Target must be `all`, `3D`, or `CSE2109`.")
        _, _, recipients, label = self.target_info(kind, val)
        status = await ctx.send(f"⏳ Sending DM to {len(recipients)} students ({label})...")
        embed = discord.Embed(title=f"📩 Message from {self.s.teacher}", description=text, color=BLUE,
                              timestamp=self.s.now())
        embed.set_footer(text=f"{ctx.guild.name} · {label}")
        sent, failed = 0, []
        for r in recipients:
            m = ctx.guild.get_member(int(r["discord_id"]))
            if not m:
                failed.append(r["name"])
                continue
            try:
                await m.send(embed=embed)
                sent += 1
            except discord.HTTPException:
                failed.append(r["name"])
            await asyncio.sleep(1.0)
        await status.edit(content=f"✅ DM sent: **{sent}** · failed: **{len(failed)}**"
                                  + (f"\nFailed: {', '.join(failed[:30])}" if failed else ""))

    # ------------------------------------------------------------ marks by DM
    @commands.command(name="marks_dm")
    async def marks_dm(self, ctx, *, title: str = "Your marks"):
        """Attach a CSV/XLSX: one row per student, a column with the Student ID, other columns = marks."""
        if not ctx.message.attachments:
            raise commands.BadArgument("Attach a CSV or Excel file (first row = column names, one column = Student ID).")
        att = ctx.message.attachments[0]
        raw = await att.read()
        rows = []
        if att.filename.lower().endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            for r in wb.active.iter_rows(values_only=True):
                rows.append(["" if v is None else str(v).strip() for v in r])
        else:
            import csv as _csv
            rows = [[c.strip() for c in r] for r in _csv.reader(io.StringIO(raw.decode("utf-8-sig", errors="replace")))]
        rows = [r for r in rows if any(r)]
        if len(rows) < 2:
            raise commands.BadArgument("The file needs a header row and at least one student row.")
        header = rows[0]
        id_re = re.compile(rf"^\d{{{self.s.id_length}}}$")
        id_col = next((i for i in range(len(header)) if any(id_re.match(r[i]) for r in rows[1:] if i < len(r))), None)
        if id_col is None:
            raise commands.BadArgument(f"No column with {self.s.id_length}-digit Student IDs found.")
        by_sid = {}
        for rec in self.bot.store.all_active():
            by_sid.setdefault(rec["student_id"], rec)
        status = await ctx.send(f"⏳ Sending marks to {len(rows) - 1} students...")
        sent, failed, unknown = 0, [], []
        for r in rows[1:]:
            sid = r[id_col] if id_col < len(r) else ""
            rec = by_sid.get(sid)
            if not rec:
                unknown.append(sid)
                continue
            m = ctx.guild.get_member(int(rec["discord_id"]))
            if not m:
                failed.append(sid)
                continue
            e = discord.Embed(title=f"📊 {title}", color=BLUE, timestamp=self.s.now(),
                              description=f"**{rec['name']}** · `{sid}` · Section {rec['section']}")
            for i, h in enumerate(header[:25]):
                if i != id_col and i < len(r) and h:
                    e.add_field(name=h[:256], value=(r[i] or "-")[:1024], inline=True)
            e.set_footer(text=f"{self.s.teacher} · private — only you can see this")
            try:
                await m.send(embed=e)
                sent += 1
            except discord.HTTPException:
                failed.append(sid)
            await asyncio.sleep(0.8)
        await status.edit(content=(f"✅ Marks sent: **{sent}** · DM closed/failed: **{len(failed)}** · "
                                   f"ID not registered: **{len(unknown)}**"
                                   + (f"\nNot registered: {', '.join(unknown[:40])}" if unknown else "")
                                   + (f"\nFailed: {', '.join(failed[:40])}" if failed else "")))

    # ------------------------------------------------------------ excel
    def _pairs(self, target: str):
        s = self.s
        parts = target.split()
        if not parts or parts[0].lower() == "all":
            return [(c, sec) for sec in s.sections for c in s.section_courses(sec)]
        secs, courses = [], []
        for p in parts:
            sec, c = s.normalize_section(p), s.normalize_course(p)
            if sec:
                secs.append(sec)
            elif c:
                courses.append(c)
            else:
                raise commands.BadArgument(f"`{p}` is not a section or course.")
        pairs = [(c, sec) for sec in (secs or list(s.sections)) for c in s.section_courses(sec)
                 if not courses or c in courses]
        if not pairs:
            raise commands.BadArgument("That course/section combination does not exist.")
        return pairs

    @commands.command(name="excel", aliases=["xlsx"])
    async def excel(self, ctx, *, target: str = "all"):
        """One Excel file: Summary sheet + one sheet per course/section."""
        pairs = self._pairs(target)
        try:
            from core.excel import build_student_workbook
            name = "students_" + ("all" if target.strip().lower() in ("", "all") else "_".join(target.split()))
            dest = DATA_DIR / "exports" / f"{name}_{self.s.now():%Y%m%d_%H%M}.xlsx"
            path = await asyncio.to_thread(build_student_workbook, self.bot.store, self.s, pairs, dest)
        except ImportError:
            raise commands.BadArgument("openpyxl is not installed yet — run `!update` (it installs requirements).")
        total = sum(len(self.bot.store.section_students(sec)) for sec in {p[1] for p in pairs})
        await ctx.send(f"📗 Excel file · {len(pairs)} sheet(s) · {total} students", file=discord.File(path))

    # ------------------------------------------------------------ export
    @commands.command(name="export")
    async def export(self, ctx, *, target: str = "all"):
        store, s = self.bot.store, self.s
        store.export_files()
        parts = target.split()
        if not parts or parts[0].lower() == "all":
            paths = sorted(STUDENTS_DIR.glob("*.csv")) + sorted(STUDENTS_DIR.glob("*.json"))
            paths.append(DATA_DIR / "verified_records.json")
            zpath = zip_files(paths, DATA_DIR / "exports" / f"students_{s.now():%Y%m%d_%H%M}.zip")
            await ctx.send("📦 All student files (CSV + JSON):", file=discord.File(zpath))
            return
        secs, courses = [], []
        for p in parts:
            sec, c = s.normalize_section(p), s.normalize_course(p)
            if sec:
                secs.append(sec)
            elif c:
                courses.append(c)
            else:
                raise commands.BadArgument(f"`{p}` is not a section or course.")
        pairs = []
        for sec in (secs or list(s.sections)):
            for c in s.section_courses(sec):
                if not courses or c in courses:
                    pairs.append((c, sec))
        if not pairs:
            raise commands.BadArgument("That course/section combination does not exist.")
        files = [discord.File(STUDENTS_DIR / f"{c}_{sec}.csv") for c, sec in pairs
                 if (STUDENTS_DIR / f"{c}_{sec}.csv").exists()][:10]
        await ctx.send(f"📄 {', '.join(f'{c}_{sec}' for c, sec in pairs)}", files=files)


async def setup(bot):
    await bot.add_cog(Admin(bot))
