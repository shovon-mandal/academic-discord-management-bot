"""Excel (.xlsx) export: one sheet per course + section, plus a summary sheet."""
from __future__ import annotations

from pathlib import Path


def build_student_workbook(store, settings, pairs, dest: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append([f"{settings.server_name}", "", "", f"Generated: {settings.now_str()}"])
    summary["A1"].font = Font(bold=True, size=13)
    summary.append([])
    summary.append(["Course Code", "Course Title", "Section", "Students"])
    for c in summary[3]:
        c.font, c.fill = head_font, head_fill

    header = ["SL", "Student ID", "Name", "Section", "Course Code", "Discord Username", "Server Nickname",
              "Status", "Verified At"]
    for course, sec in pairs:
        info = settings.courses[course]
        rows = store.section_students(sec)  # active students only
        summary.append([info["code"], info["title"], sec + (" (Evening)" if settings.is_evening(sec) else ""), len(rows)])
        ws = wb.create_sheet(title=f"{course}_{sec}"[:31])
        ws.append([f"{info['code']} — {info['title']} · Section {sec} · {settings.semester}"])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        ws.append(header)
        for c in ws[3]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center")
        for i, r in enumerate(rows, 1):
            ws.append([i, r.get("student_id"), r.get("name"), sec, info["code"], r.get("discord_name"),
                       r.get("nickname"), r.get("status", "active"), r.get("verified_at")])
        for row in ws.iter_rows(min_row=4, min_col=2, max_col=2):
            for c in row:
                c.number_format = "@"  # keep Student ID as text
        widths = [5, 14, 30, 9, 11, 22, 30, 9, 20]
        for idx, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = w
        ws.freeze_panes = "A4"
    for idx, w in enumerate([13, 38, 14, 10], 1):
        summary.column_dimensions[get_column_letter(idx)].width = w
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest
